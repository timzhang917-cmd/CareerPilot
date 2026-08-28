from pathlib import Path
from docx import Document
from pypdf import PdfReader

def extract_text_from_docx(file_path): # 提取word文本信息
    document = Document(file_path) # 打开文件后，返回Document对象
    text_parts = [] # 空列表，逐步收集文字

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            text_parts.append(text)

    for table in document.tables: # 要考虑到有些简历使用的内嵌表格布局
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if cells:
                text_parts.append(" | ".join(cells))

    return "\n".join(text_parts)


def extract_text_from_pdf(file_path):
    reader = PdfReader(file_path)
    text_parts = []

    for page in reader.pages:
        text = page.extract_text() or ""

        if text.strip():
            text_parts.append(text.strip())

    return "\n".join(text_parts)


def extract_resume_text(file_path):  # 分发函数
    extension = Path(file_path).suffix.lower()

    if extension == ".docx":
        return extract_text_from_docx(file_path)

    if extension == ".pdf":
        return extract_text_from_pdf(file_path)